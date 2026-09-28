"""Файл «Загрузка пользователей» для LMS: оплатившие частные лица, готовые к зачислению.

Формат — шаблон кейсодержателя один в один: 30 колонок на первом листе (заголовки дословно,
как в шаблоне, — загрузчик LMS сверяет их по тексту), списки «Пол» и «Образование» на втором
листе и выпадающие списки в колонках L и W. CRM хранит только ФИО, телефон и почту; паспорт,
СНИЛС и диплом она не собирает, поэтому эти колонки остаются пустыми — их заполняет LMS.
"""

import io
import re
import uuid
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import crypto
from app.core.scope import apply_interaction_scope
from app.core.security import CurrentUser
from app.modules.audit.models import AuditLog
from app.modules.clients.models import Client
from app.modules.interactions.models import Interaction
from app.modules.payments.models import Payment
from app.modules.workflow.models import Stage

HEADERS = (
    "Фамилия",
    "Имя",
    "Отчествопри наличии)",
    "Номер телефона",
    "Email",
    "СНИЛС",
    "Серия паспорта",
    "Номер паспорта",
    "Кем выдан паспорт",
    "Дата выдачи паспорта",
    "Код подразделения",
    "Пол",
    "Дата рождения",
    "Регион регистрации",
    "Населенный пункт регистрации",
    "Улица регистрации",
    "Дом регистрации",
    "Квартира регистрации",
    "Индекс регистрации",
    "Имядательный падеж)",
    "Фамилиядательный падеж)",
    "Отчестводательный падеж)",
    "Образование",
    "Профессия по диплому",
    "Учебное заведение по диплому",
    "Фамилия, указанная в дипломе",
    "Номер диплома",
    "Серия диплома",
    "Регистрационный номер диплома",
    "Дата выдачи диплома",
)
SEXES = ("М", "Ж")
EDUCATION = (
    "Без образования",
    "Основное общее образование - 9 классов",
    "Среднее общее образование - 11 классов",
    "Среднее профессиональное образование",
    "Высшее образование – бакалавриат",
    "Высшее образование – специалитет, магистратура",
    "Высшее образование – подготовка кадров высшей квалификации",
)
# Выпадающие списки шаблона: столько строк, сколько размечено в нём.
VALIDATED_ROWS = 1001
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def split_name(name: str) -> tuple[str, str, str]:
    """«Черепанова Светлана Васильевна» → фамилия, имя, отчество; отчества может не быть."""
    parts = name.split()
    if not parts:
        return "", "", ""
    return parts[0], parts[1] if len(parts) > 1 else "", " ".join(parts[2:])


def phone_number(phone: str | None) -> int | None:
    """В шаблоне телефон — число 79990234365, без плюса и скобок."""
    digits = re.sub(r"\D", "", phone or "")
    return int(digits) if len(digits) == 11 else None


def render(rows: list[tuple[str, str, str, int | None, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for last, first, middle, phone, email in rows:
        sheet.append([last, first, middle or None, phone, email])
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24

    lists = workbook.create_sheet("Лист2")
    for index, sex in enumerate(SEXES, start=1):
        lists.cell(row=index, column=1, value=sex)
    for index, level in enumerate(EDUCATION, start=1):
        lists.cell(row=index, column=2, value=level)

    sexes = DataValidation(type="list", formula1="Лист2!$A$1:$A$2", allow_blank=True)
    sexes.add(f"L1:L{VALIDATED_ROWS}")
    education = DataValidation(type="list", formula1="Лист2!$B$1:$B$7", allow_blank=True)
    education.add(f"W1:W{VALIDATED_ROWS}")
    sheet.add_data_validation(sexes)
    sheet.add_data_validation(education)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def export_lms_users(
    session: AsyncSession,
    user: CurrentUser,
    stage_code: str = "enrollment",
    program_id: uuid.UUID | None = None,
    stream: int | None = None,
    trace_id: str | None = None,
) -> tuple[bytes, str]:
    """Частные лица на этапе зачисления в области видимости пользователя, по курсу и потоку."""
    stmt = (
        select(Client, Interaction.id)
        .join(Interaction, Interaction.client_id == Client.id)
        .join(Stage, Stage.id == Interaction.current_stage_id)
        .where(
            Client.kind == "person",
            Client.archived_at.is_(None),
            Stage.code == stage_code,
            Interaction.status == "active",
        )
        .order_by(Client.name)
    )
    if program_id is not None:
        stmt = stmt.where(Interaction.program_id == program_id)
    if stream is not None:
        stmt = stmt.where(
            Interaction.id.in_(select(Payment.interaction_id).where(Payment.stream_no == stream))
        )
    found = (await session.execute(apply_interaction_scope(stmt, user))).tuples().all()

    rows: list[tuple[str, str, str, int | None, str]] = []
    seen: set[uuid.UUID] = set()
    for client, _interaction_id in found:
        # Человек, оплативший два курса, в файле зачисления одной строкой.
        if client.id in seen:
            continue
        seen.add(client.id)
        last, first, middle = split_name(client.name)
        rows.append(
            (
                last,
                first,
                middle,
                phone_number(crypto.decrypt(client.phone_enc)),
                crypto.decrypt(client.email_enc) or "",
            )
        )
    # В файле почта и телефоны людей: выгрузка — обращение к персональным данным.
    session.add(
        AuditLog(
            actor_user_id=user.id,
            action="client.exported_to_lms",
            entity_kind="client",
            after={
                "rows": len(rows),
                "stage": stage_code,
                "program_id": str(program_id) if program_id else None,
                "stream": stream,
            },
            trace_id=trace_id,
        )
    )
    await session.commit()
    stamp = datetime.now(UTC).strftime("%Y-%m-%d")
    return render(rows), f"Загрузка пользователей-{stamp}.xlsx"
