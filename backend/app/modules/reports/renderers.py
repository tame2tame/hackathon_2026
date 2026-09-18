"""Отчёт в файл: xlsx, xls, csv, json и pdf. PDF строится без браузера (ADR-013).

Табличные форматы открываются в Excel, поэтому значения, начинающиеся с `=`, `+`, `-` или `@`,
экранируются апострофом: название вуза «=HYPERLINK(...)» не должно стать формулой у получателя.
"""

import csv
import io
import json
from pathlib import Path

import xlwt
from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook

MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "csv": "text/csv",
    "pdf": "application/pdf",
    "json": "application/json",
}
# Кодировка CSV → кодек Python. UTF-8 пишется с BOM: без него Excel открывает файл как cp1251.
CSV_CODECS = {"utf-8": "utf-8-sig", "windows-1251": "cp1251"}
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: str) -> str:
    return f"'{value}" if value.startswith(FORMULA_START) else value


# Шрифт с кириллицей: в образе — пакет fonts-dejavu-core, на macOS — системный Arial Unicode.
FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
)
PDF_FONT_SIZE = 7
PDF_MARGIN_MM = 8


class FontMissingError(RuntimeError):
    """Без шрифта с кириллицей PDF получился бы нечитаемым, поэтому отчёт не строится."""


def find_font() -> Path:
    for candidate in FONT_CANDIDATES:
        path = Path(candidate)
        if path.exists():
            return path
    raise FontMissingError(
        "Не найден шрифт с кириллицей: установите fonts-dejavu-core или укажите свой."
    )


def render(
    fmt: str, headers: list[str], rows: list[list[str]], title: str, encoding: str = "utf-8"
) -> bytes:
    match fmt:
        case "xlsx":
            return _xlsx(headers, rows, title)
        case "xls":
            return _xls(headers, rows, title)
        case "csv":
            return _csv(headers, rows, encoding)
        case "json":
            return _json(headers, rows)
        case _:
            return _pdf(headers, rows, title)


def _csv(headers: list[str], rows: list[list[str]], encoding: str) -> bytes:
    buffer = io.StringIO(newline="")
    # Точка с запятой — разделитель Excel в русской локали; строка с переносом берётся в кавычки.
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([safe_cell(header) for header in headers])
    writer.writerows([[safe_cell(value) for value in row] for row in rows])
    # Символ, которого нет в cp1251, заменяется «?», а не роняет отчёт; без потерь — UTF-8.
    return buffer.getvalue().encode(CSV_CODECS[encoding], errors="replace")


def _xlsx(headers: list[str], rows: list[list[str]], title: str) -> bytes:
    # write_only: строки уходят в файл по мере записи, не собираясь в памяти целиком.
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title[:31])
    sheet.append([safe_cell(header) for header in headers])
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xls(headers: list[str], rows: list[list[str]], title: str) -> bytes:
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet(title[:31])
    bold = xlwt.easyxf("font: bold on")
    for column, header in enumerate(headers):
        sheet.write(0, column, safe_cell(header), bold)
    for row_no, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            sheet.write(row_no, column, safe_cell(value))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _json(headers: list[str], rows: list[list[str]]) -> bytes:
    items = [dict(zip(headers, row, strict=True)) for row in rows]
    return json.dumps(items, ensure_ascii=False, indent=2).encode("utf-8")


def _pdf(headers: list[str], rows: list[list[str]], title: str) -> bytes:
    font = find_font()
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_margin(PDF_MARGIN_MM)
    pdf.add_font("report", "", str(font))
    pdf.set_font("report", size=PDF_FONT_SIZE + 3)
    pdf.add_page()
    pdf.cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("report", size=PDF_FONT_SIZE)

    width = (pdf.w - 2 * PDF_MARGIN_MM) / max(1, len(headers))
    # Шапка без жирного начертания: подключён один файл шрифта, а выделяет её заливка.
    headings = FontFace(emphasis="", fill_color=(235, 235, 235))
    with pdf.table(
        col_widths=tuple([width] * len(headers)), line_height=4, headings_style=headings
    ) as table:
        head = table.row()
        for header in headers:
            head.cell(header)
        for row in rows:
            line = table.row()
            for value in row:
                line.cell(value)
    return bytes(pdf.output())
