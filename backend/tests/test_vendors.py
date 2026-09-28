"""Вендоры с контактами: таблица кейсодержателя «Вендоры» загружается как есть."""

import io
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.audit.models import AuditLog
from app.modules.vendors.models import VendorContact
from tests.users import ALINA_ADMIN, ANNA_KAM, ROMAN_MANAGER, as_user

IMPORT = "/api/v1/admin/catalogs/vendor-contacts/import"
HEADERS = ["Компания", "Продукт", "ФИО", "Телефон", "Почта", "Способ связи"]
# Строки в том виде, в каком их прислал кейсодержатель (люди вымышленные, домен example.ru).
ROWS = [
    ["ООО «Базис»", "«Базис Dynamix»", "Иванов Иван Иванович", "+7 (900) 111-22-33",
     "ivanov.ii@example.ru", "Почта, Чат в ТГ"],
    ["ООО «ТДата»", "«RT.DataLake», «RT.Warehouse»", "Смирнова Анна Петровна",
     "+7 (911) 222-33-44", "smirnova.ap@example.ru", "Чат в ТГ"],
    ["ООО «РТК ИТ»", "«Нейрошлюз»", "Новикова Ольга Александровна", "+7 (977) 888-99-00",
     "novikova.oa@example.ru", "Почта"],
]  # fmt: skip


def workbook(rows: list[list[str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def load(client: AsyncClient, rows: list[list[str]], *, dry_run: bool = False) -> Any:
    response = await client.post(
        IMPORT,
        files={"file": ("Вендоры.xlsx", workbook(rows), "application/octet-stream")},
        data={"dry_run": str(dry_run).lower()},
        headers=as_user(ALINA_ADMIN),
    )
    assert response.status_code == 200, response.text
    return response.json()


async def vendor_named(client: AsyncClient, name: str) -> dict[str, Any]:
    vendors = (await client.get("/api/v1/vendors", headers=as_user(ANNA_KAM))).json()
    return next(item for item in vendors if item["name"] == name)


async def test_case_table_loads_as_is_and_twice_changes_nothing(
    client: AsyncClient, session: AsyncSession
) -> None:
    first = await load(client, ROWS)
    again = await load(client, ROWS)

    assert first["created"] == 3, first
    assert again["unchanged"] == 3, again
    tdata = await vendor_named(client, "ООО «ТДата»")
    # Одна строка — два продукта в кавычках: оба заведены и оба за одним человеком.
    assert sorted(product["name"] for product in tdata["products"]) == [
        "RT.DataLake",
        "RT.Warehouse",
    ]
    assert tdata["contacts"] == 1

    contacts = await client.get(
        f"/api/v1/vendors/{tdata['id']}/contacts", headers=as_user(ANNA_KAM)
    )
    [contact] = contacts.json()
    assert contact["full_name"] == "Смирнова Анна Петровна"
    assert (contact["email"], contact["phone"]) == ("smirnova.ap@example.ru", "+7 (911) 222-33-44")
    assert contact["channels"] == ["telegram"]
    assert len(contact["products"]) == 2
    # Почта хранится зашифрованной, а просмотр контактов пишется в журнал.
    stored = await session.scalar(
        select(VendorContact).where(VendorContact.full_name.like("Смир%"))
    )
    assert stored is not None
    assert stored.email_enc is not None
    assert b"example.ru" not in stored.email_enc
    viewed = await session.scalar(
        select(AuditLog).where(AuditLog.action == "vendor_contact.viewed")
    )
    assert viewed is not None


async def test_bad_rows_are_reported_one_by_one(client: AsyncClient) -> None:
    result = await load(
        client,
        [
            ["ООО «Базис»", "«Базис Dynamix»", "Иванов Иван Иванович", "", "ivanov", "Почта"],
            ["ООО «Базис»", "«Базис Dynamix»", "Петров Пётр", "", "", "Голубиная почта"],
            ROWS[2],
        ],
    )

    actions = [(row["action"], row["detail"]) for row in result["rows"]]
    assert actions[0] == ("error", "Это не почта: «ivanov»")
    assert actions[1] == ("error", "Не понятный способ связи: «Голубиная почта»")
    assert actions[2][0] == "created"


async def test_export_is_the_same_table_and_audited(
    client: AsyncClient, session: AsyncSession
) -> None:
    await load(client, ROWS)

    exported = await client.get(
        "/api/v1/admin/catalogs/vendor-contacts/export", headers=as_user(ROMAN_MANAGER)
    )

    assert exported.status_code == 200, exported.text
    rows = list(load_workbook(io.BytesIO(exported.content)).active.values)
    assert list(rows[0]) == HEADERS
    # Телефон с «+» выгружается с апострофом — защита от формул в Excel (ADR-019);
    # загрузка апостроф снимает, поэтому файл возвращается без изменений.
    assert (
        "ООО «ТДата»",
        "«RT.DataLake», «RT.Warehouse»",
        "Смирнова Анна Петровна",
        "'+7 (911) 222-33-44",
        "smirnova.ap@example.ru",
        "Чат в ТГ",
    ) in rows
    audited = await session.scalar(
        select(AuditLog).where(AuditLog.action == "vendor_contact.exported")
    )
    assert audited is not None
    reloaded = await load(client, [list(row) for row in rows[1:]])
    assert reloaded["unchanged"] == 3


async def test_contacts_are_kept_by_managers(client: AsyncClient) -> None:
    await load(client, ROWS[:2])
    basis = await vendor_named(client, "ООО «Базис»")
    tdata = await vendor_named(client, "ООО «ТДата»")
    body = {"full_name": "Орлова Ника", "email": "n.orlova@example.ru", "channels": ["email"]}

    by_kam = await client.post(
        f"/api/v1/vendors/{basis['id']}/contacts", json=body, headers=as_user(ANNA_KAM)
    )
    foreign_product = await client.post(
        f"/api/v1/vendors/{basis['id']}/contacts",
        json={**body, "product_ids": [tdata["products"][0]["id"]]},
        headers=as_user(ROMAN_MANAGER),
    )
    created = await client.post(
        f"/api/v1/vendors/{basis['id']}/contacts",
        json={**body, "product_ids": [basis["products"][0]["id"]]},
        headers=as_user(ROMAN_MANAGER),
    )
    archived = await client.post(
        f"/api/v1/vendor-contacts/{created.json()['id']}/archive",
        headers=as_user(ROMAN_MANAGER),
    )

    assert by_kam.status_code == 403
    assert foreign_product.status_code == 422
    assert created.status_code == 201, created.text
    assert created.json()["products"] == [basis["products"][0]]
    # В архиве человеку незачем оставлять почту и телефон.
    assert (archived.json()["email"], archived.json()["phone"]) == (None, None)
