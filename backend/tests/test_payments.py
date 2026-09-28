"""Оплаты частных лиц: выгрузка платёжной системы → клиент, запись B2C, файл зачисления в LMS.

Структура файла — как у кейсодержателя (включая `null` среди записей), люди синтетические.
"""

import io
import json
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.audit.models import AuditLog
from app.modules.catalogs.models import Direction, Program
from app.modules.clients.models import Client
from app.modules.interactions.models import Interaction, Transition
from app.modules.payments import lms
from app.modules.payments.models import Payment
from app.modules.workflow.models import Stage
from tests.test_groups import b2c_record, catalog_id
from tests.users import ALINA_ADMIN, ANNA_KAM, ROMAN_MANAGER, as_user

IMPORT = "/api/v1/payments/import"
LMS_USERS = "/api/v1/payments/lms-users"


def payment(order: str, course: str = "Промпт-инжиниринг", **fields: Any) -> dict[str, Any]:
    return {
        "Номер заявки": order,
        "Курс": course,
        "Фамилия": "Тестова",
        "Имя": "Анна",
        "Отчество": "Олеговна",
        "Телефон": "7 (999) 000-11-22",
        "Email": "a.testova@example.com",
        "Номер потока": 1,
        **fields,
    }


async def prompt_course(session: AsyncSession) -> Program:
    """Продуктонезависимый курс из выгрузки кейсодержателя."""
    direction = await session.scalar(select(Direction).where(Direction.code == "ai"))
    assert direction is not None
    program = Program(direction_id=direction.id, name="Промпт-инжиниринг")
    session.add(program)
    await session.commit()
    return program


async def upload(
    client: AsyncClient, items: list[Any], *, dry_run: bool, user: str = ROMAN_MANAGER
) -> Any:
    content = json.dumps(items, ensure_ascii=False).encode()
    return await client.post(
        IMPORT,
        files={"file": ("Данные оплат.json", content, "application/json")},
        data={"dry_run": str(dry_run).lower()},
        headers=as_user(user),
    )


async def stage_code(session: AsyncSession, interaction_id: Any) -> str:
    session.expire_all()
    interaction = await session.get(Interaction, interaction_id)
    assert interaction is not None
    stage = await session.get(Stage, interaction.current_stage_id)
    assert stage is not None
    return stage.code


async def test_payment_creates_client_and_record_ready_for_enrollment(
    client: AsyncClient, session: AsyncSession
) -> None:
    await prompt_course(session)
    items = [
        None,
        payment("ORD-1"),
        payment("ORD-2", course="Курс, которого нет"),
        payment("ORD-3", **{"Email": "без собаки"}),
    ]

    preview = await upload(client, items, dry_run=True)
    nothing = await session.scalar(select(Payment.id))
    applied = await upload(client, items, dry_run=False)

    assert preview.status_code == 200, preview.text
    assert nothing is None
    body = applied.json()
    assert (body["created"], body["errors"]) == (1, 3)
    details = [row["detail"] for row in body["rows"]]
    assert details[0] == "Пустая запись"
    assert details[2] == "Курса «Курс, которого нет» нет в справочнике программ"
    assert details[3] == "Это не почта: «без собаки»"

    stored = await session.scalar(select(Payment).where(Payment.order_no == "ORD-1"))
    assert stored is not None
    assert stored.stream_no == 1
    person = await session.get(Client, stored.client_id)
    assert person is not None
    assert (person.kind, person.name) == ("person", "Тестова Анна Олеговна")
    # Почта и телефон — зашифрованы, как у любой карточки клиента.
    assert person.email_enc is not None
    assert b"example.com" not in person.email_enc
    interaction_id = stored.interaction_id
    assert await stage_code(session, interaction_id) == "enrollment"
    history = await session.scalar(
        select(Transition).where(Transition.interaction_id == interaction_id)
    )
    assert history is not None
    assert history.comment == "Оплата по заявке ORD-1, поток 1"
    # В журнал — номер заявки и поток, но не почта и не ФИО.
    entries = list(await session.scalars(select(AuditLog).where(AuditLog.action.like("payment.%"))))
    assert entries
    assert all(
        "example.com" not in json.dumps(entry.after, ensure_ascii=False) for entry in entries
    )
    assert all("Тестова" not in json.dumps(entry.after, ensure_ascii=False) for entry in entries)


async def test_same_file_twice_changes_nothing(client: AsyncClient, session: AsyncSession) -> None:
    await prompt_course(session)
    await upload(client, [payment("ORD-1")], dry_run=False)

    again = await upload(client, [payment("ORD-1")], dry_run=False)
    second_course_payment = await upload(
        client, [payment("ORD-9", **{"Номер потока": 2})], dry_run=False
    )

    assert again.json()["unchanged"] == 1
    # Тот же человек и тот же курс: новая оплата не заводит вторую карточку и вторую запись.
    assert second_course_payment.json()["updated"] == 1
    assert len(list(await session.scalars(select(Client).where(Client.kind == "person")))) == 1
    payments = list(await session.scalars(select(Payment)))
    assert len({item.interaction_id for item in payments}) == 1


async def test_payment_moves_a_known_record_to_enrollment(
    client: AsyncClient, session: AsyncSession
) -> None:
    # Руководитель уже ведёт этого человека: заявка на DevOps на первом этапе.
    record = await b2c_record(client, ROMAN_MANAGER)
    paid = payment(
        "ORD-7",
        course="DevOps-инженерия",
        **{"Фамилия": "Петров", "Имя": "Иван", "Отчество": "", "Email": "Ivan.Petrov@example.com"},
    )

    response = await upload(client, [paid], dry_run=False)

    assert response.json()["rows"][0]["detail"] == "переведена на «Зачисление в LMS»"
    assert await stage_code(session, record["id"]) == "enrollment"
    persons = list(await session.scalars(select(Client).where(Client.kind == "person")))
    assert len(persons) == 1


async def test_only_managers_load_payments(client: AsyncClient, session: AsyncSession) -> None:
    await prompt_course(session)

    by_kam = await upload(client, [payment("ORD-1")], dry_run=True, user=ANNA_KAM)
    by_admin = await upload(client, [payment("ORD-1")], dry_run=True, user=ALINA_ADMIN)

    assert by_kam.status_code == 403
    assert by_admin.status_code == 200


async def test_phone_formats_are_one_number() -> None:
    from app.modules.payments.service import phone_of

    assert phone_of("7 (999) 023-43-65") == "+7 (999) 023-43-65"
    assert phone_of("89990234365") == "+7 (999) 023-43-65"
    assert phone_of("9990234365") == "+7 (999) 023-43-65"
    assert phone_of("") is None


async def test_lms_file_follows_the_template(client: AsyncClient, session: AsyncSession) -> None:
    await prompt_course(session)
    await upload(
        client,
        [
            payment("ORD-1"),
            payment(
                "ORD-2",
                **{"Фамилия": "Второва", "Email": "vtorova@example.com", "Номер потока": 2},
            ),
        ],
        dry_run=False,
    )
    program_id = await catalog_id(client, "programs", "Промпт-инжиниринг")

    everyone = await client.get(
        LMS_USERS, params={"program_id": program_id}, headers=as_user(ROMAN_MANAGER)
    )
    second_stream = await client.get(
        LMS_USERS, params={"stream": 2}, headers=as_user(ROMAN_MANAGER)
    )
    foreign = await client.get(LMS_USERS, headers=as_user(ANNA_KAM))

    assert everyone.status_code == 200, everyone.text
    workbook = load_workbook(io.BytesIO(everyone.content))
    sheet = workbook["Лист1"]
    rows = list(sheet.iter_rows(values_only=True))
    # Заголовки дословно из шаблона кейсодержателя, строки — по алфавиту фамилий.
    assert rows[0] == lms.HEADERS
    assert rows[1][:5] == ("Второва", "Анна", "Олеговна", 79990001122, "vtorova@example.com")
    assert rows[2][:5] == ("Тестова", "Анна", "Олеговна", 79990001122, "a.testova@example.com")
    assert all(value is None for value in rows[1][5:])
    lists = workbook["Лист2"]
    assert [lists.cell(row=i, column=1).value for i in (1, 2)] == ["М", "Ж"]
    assert lists.cell(row=7, column=2).value == lms.EDUCATION[-1]
    ranges = {str(item.sqref): item.formula1 for item in sheet.data_validations.dataValidation}
    assert ranges == {"L1:L1001": "Лист2!$A$1:$A$2", "W1:W1001": "Лист2!$B$1:$B$7"}

    streamed = list(load_workbook(io.BytesIO(second_stream.content))["Лист1"].values)
    assert [row[0] for row in streamed[1:]] == ["Второва"]
    # Записи ведёт руководитель: у КАМа без доступа к ним файл пустой, а не чужой.
    assert list(load_workbook(io.BytesIO(foreign.content))["Лист1"].values)[1:] == []
    exported = await session.scalar(
        select(AuditLog).where(AuditLog.action == "client.exported_to_lms")
    )
    assert exported is not None
